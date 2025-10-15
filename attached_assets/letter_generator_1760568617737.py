"""
Cover Letter Generator - Create personalized cover letters using AI
"""

import logging
import os
from typing import Dict, Any
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from src.ai_client import get_ai_client

logger = logging.getLogger(__name__)

LETTERS_DIR = 'data/generated/cover_letters'
os.makedirs(LETTERS_DIR, exist_ok=True)

def generate_cover_letter(job_info: Dict[str, Any], cv_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Generate personalized cover letter
    
    Args:
        job_info: Job details (title, company, description, requirements)
        cv_data: Parsed CV data
    
    Returns:
        Dict with 'text' and optionally 'docx_path'
    """
    try:
        ai_client = get_ai_client()
        letter_text = ai_client.generate_cover_letter(job_info, cv_data)
        
        # Generate filename
        company = job_info.get('company', 'Company').replace(' ', '_')
        title = job_info.get('title', 'Position').replace(' ', '_')
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f"CoverLetter_{company}_{title}_{timestamp}"
        
        # Save as DOCX
        docx_path = os.path.join(LETTERS_DIR, f"{filename}.docx")
        doc = create_letter_docx(letter_text, cv_data)
        doc.save(docx_path)
        
        # Save as TXT
        txt_path = os.path.join(LETTERS_DIR, f"{filename}.txt")
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write(letter_text)
        
        logger.info(f"Generated cover letter: {filename}")
        
        return {
            'text': letter_text,
            'docx_path': docx_path,
            'txt_path': txt_path,
            'filename': filename
        }
        
    except Exception as e:
        logger.error(f"Cover letter generation error: {e}")
        return {
            'error': str(e),
            'text': f"Error generating cover letter: {e}"
        }

def create_letter_docx(letter_text: str, cv_data: Dict[str, Any]) -> Document:
    """Create formatted DOCX document for cover letter"""
    doc = Document()
    
    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    # Add header with applicant info
    header = doc.add_paragraph()
    header_run = header.add_run(f"{cv_data.get('name', 'Applicant')}\n")
    header_run.bold = True
    header_run.font.size = Pt(14)
    
    contact_run = header.add_run(
        f"{cv_data.get('email', '')}\n"
        f"{cv_data.get('phone', '')}\n"
        f"{cv_data.get('location', '')}"
    )
    contact_run.font.size = Pt(10)
    
    # Add date
    doc.add_paragraph()
    date_para = doc.add_paragraph(datetime.now().strftime('%B %d, %Y'))
    date_para.alignment = 0  # Left align
    
    # Add letter body
    doc.add_paragraph()
    
    # Split letter into paragraphs
    paragraphs = letter_text.split('\n\n')
    for para_text in paragraphs:
        if para_text.strip():
            para = doc.add_paragraph(para_text.strip())
            para.alignment = 3  # Justify
            para_format = para.paragraph_format
            para_format.space_after = Pt(12)
    
    return doc

def get_sample_letter() -> str:
    """Get a sample cover letter for demo purposes"""
    return """Dear Hiring Manager,

I am writing to express my strong interest in the Software Engineer position at TechCorp. With over 5 years of experience building scalable web applications and a proven track record of delivering high-quality solutions, I am confident I would be a valuable addition to your team.

In my current role as Senior Software Engineer, I have successfully led the development of multiple full-stack applications using Python, JavaScript, and modern frameworks like React and Node.js. I am particularly drawn to TechCorp's commitment to innovation and would be excited to contribute my expertise in cloud architecture and DevOps practices to your projects.

My technical skills align well with your requirements, including proficiency in Python, React, SQL, Docker, and AWS. Beyond technical abilities, I bring strong problem-solving skills, attention to detail, and a collaborative mindset that has enabled me to thrive in fast-paced startup environments.

I would welcome the opportunity to discuss how my background and skills would benefit TechCorp. Thank you for your consideration.

Sincerely,
Demo User"""
